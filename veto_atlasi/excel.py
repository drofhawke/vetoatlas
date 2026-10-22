"""
Açık veri CSV'lerinin Excel kopyası (web/veri/*.xlsx).

Türkçe bölge ayarlı Excel çift tıklamada CSV'yi ";" ile ayırır; atlasın CSV'leri standart virgüllüdür. Bu yüzden sütunlar
bölünmüyordu (14 Ekim 2026 denemesi). XLSX kopyası aynı satırları taşır ve her ayarda doğru açılır.

Çıktı baytları belirlenimlidir (04 her çalıştığında aynı dosya): belge özelliklerindeki tarihler sabit, zip girdileri sabit
tarihli ve sıkıştırmasız (sıkıştırma baytları zlib sürümüne bağlıdır). Hücreler: tam sayılar sayı, geri kalanı metin
(tarihler metin kalır, Excel bölge ayarına göre dönüştürmez; "=" ile başlayan metin formül sayılmaz).
"""
from __future__ import annotations

import csv
import datetime
import io
import re
import zipfile

from openpyxl import Workbook
from openpyxl.styles import Font

SABIT_TARIH = datetime.datetime(2000, 1, 1)
_TAM_SAYI = re.compile(r"^-?(?:0|[1-9]\d{0,14})$")


def _hucre(deger: str) -> int | str | None:
    if deger == "":
        return None
    return int(deger) if _TAM_SAYI.match(deger) else deger


def csv_satirlari(veri: bytes) -> list[list[str]]:
    """UTF-8 CSV baytlarından satırlar (BOM'lu ya da BOM'suz)."""
    return list(csv.reader(io.StringIO(veri.decode("utf-8-sig"), newline="")))


def _sayfa_doldur(ws, satirlar: list[list[str]]) -> None:
    for satir in satirlar:
        ws.append([_hucre(x) for x in satir])
    for satir in ws.iter_rows():
        for h in satir:
            if isinstance(h.value, str):
                h.data_type = "s"  # "=" ile başlayan metin de metin kalır
    for h in ws[1]:
        h.font = Font(bold=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def xlsx(satirlar: list[list[str]], sayfa: str, sozluk: list[list[str]] | None = None) -> bytes:
    """
    Satırlardan belirlenimli bir XLSX: ilk satır kalın başlık, başlık satırı sabit, süzgeç açık.
    sozluk verilirse (başlık satırıyla birlikte) ikinci sayfa "sozluk" olarak eklenir; ilk sayfa etkin kalır.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = sayfa
    _sayfa_doldur(ws, satirlar)
    if sozluk:
        _sayfa_doldur(wb.create_sheet("sozluk"), sozluk)
    wb.properties.creator = "Veto Atlası"
    wb.properties.created = wb.properties.modified = SABIT_TARIH
    ham = io.BytesIO()
    wb.save(ham)
    cikti = io.BytesIO()
    with zipfile.ZipFile(ham) as z, zipfile.ZipFile(cikti, "w", zipfile.ZIP_STORED) as y:
        for bilgi in z.infolist():
            yeni = zipfile.ZipInfo(bilgi.filename, date_time=(1980, 1, 1, 0, 0, 0))
            yeni.external_attr = 0o600 << 16
            veri = z.read(bilgi.filename)
            if bilgi.filename == "docProps/core.xml":
                # openpyxl kaydederken "modified" tarihini o anın saatiyle yazar; sabit tarihe çekilir
                veri = re.sub(rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)",
                              rb"\g<1>" + SABIT_TARIH.strftime("%Y-%m-%dT%H:%M:%SZ").encode() + rb"\g<2>", veri)
            y.writestr(yeni, veri)
    return cikti.getvalue()
