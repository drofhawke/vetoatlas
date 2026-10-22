"""
Açık veri dosyalarının Excel'de doğru açılması (14 Ekim 2026 denemesi, Türkçe bölge ayarlı Excel 16, çift tıklama):
BOM'suz CSV'de Türkçe karakterler bozuktu; virgüllü CSV'lerde sütunlar bölünmüyordu (Excel ";" bekler). Çözüm: iki CSV de
UTF-8 BOM'lu ve her biri için aynı satırları taşıyan belirlenimli bir XLSX kopyası (veto_atlasi/excel.py, 04 üretir).
"""
import csv
import io

from openpyxl import load_workbook

from conftest import KOK
from veto_atlasi import excel

VERI = KOK / "web/veri"
ADLAR = ("vetolar", "genel_kurul_tartismalari")


def test_iki_csv_de_bom_ile_basliyor():
    for ad in ADLAR:
        assert (VERI / f"{ad}.csv").read_bytes().startswith(b"\xef\xbb\xbf"), ad
    # Kaynak dosya BOM'suz kalır; açık veri kopyası kaynağın başına BOM eklenmiş hâlidir
    kaynak = (KOK / "kaynaklar/genel_kurul_tartismalari.csv").read_bytes()
    assert not kaynak.startswith(b"\xef\xbb\xbf")
    assert (VERI / "genel_kurul_tartismalari.csv").read_bytes() == b"\xef\xbb\xbf" + kaynak


def test_xlsx_csv_ile_hucre_hucre_ayni():
    for ad in ADLAR:
        satirlar = list(csv.reader(io.StringIO((VERI / f"{ad}.csv").read_text(encoding="utf-8-sig"), newline="")))
        ws = load_workbook(VERI / f"{ad}.xlsx").active
        assert ws.title == ad and ws.max_row == len(satirlar) and ws.max_column == len(satirlar[0]), ad
        for i, satir in enumerate(satirlar, 1):
            degerler = ["" if h.value is None else str(h.value) for h in ws[i]]
            assert degerler == satir, (ad, i)
        assert ws.freeze_panes == "A2" and all(h.font.b for h in ws[1])


def _sozluk(ad: str) -> list[list[str]]:
    """Veri sözlüğünün başlığı ve bu dosyanın satırları (04'ün XLSX'e "sozluk" sayfası olarak eklediği)."""
    satirlar = excel.csv_satirlari((VERI / "veri_sozlugu.csv").read_bytes())
    i = satirlar[0].index("dosya")
    return [satirlar[0]] + [s for s in satirlar[1:] if s[i] == f"{ad}.csv"]


def test_xlsx_belirlenimli_ve_guncel():
    """04 her çalıştığında aynı baytlar: diskteki XLSX, CSV'den ve veri sözlüğünden yeniden üretilenle bayt bayt aynı."""
    for ad in ADLAR:
        yeni = excel.xlsx(excel.csv_satirlari((VERI / f"{ad}.csv").read_bytes()), ad, _sozluk(ad))
        assert yeni == (VERI / f"{ad}.xlsx").read_bytes(), ad
        assert yeni == excel.xlsx(excel.csv_satirlari((VERI / f"{ad}.csv").read_bytes()), ad, _sozluk(ad))


def test_veri_sozlugu_csv_basliklariyla_birebir_ve_xlsx_ikinci_sayfasi():
    """K10: veri sözlüğü her açık veri CSV'sinin sütunlarını sırasıyla tanımlar; açık veri kopyası kaynağın BOM'lu hâli;
    XLSX'lerin ikinci sayfası "sozluk" bu satırları taşır, ilk sayfa etkin kalır."""
    kaynak = (KOK / "kaynaklar/veri_sozlugu.csv").read_bytes()
    assert not kaynak.startswith(b"\xef\xbb\xbf") and (VERI / "veri_sozlugu.csv").read_bytes() == b"\xef\xbb\xbf" + kaynak
    for ad in ADLAR:
        sozluk = _sozluk(ad)
        i = sozluk[0].index("sutun")
        assert [s[i] for s in sozluk[1:]] == excel.csv_satirlari((VERI / f"{ad}.csv").read_bytes())[0], ad
        wb = load_workbook(VERI / f"{ad}.xlsx")
        assert wb.sheetnames == [ad, "sozluk"] and wb.active.title == ad
        assert [["" if h.value is None else str(h.value) for h in r] for r in wb["sozluk"].iter_rows()] == sozluk, ad


def test_hucre_turleri_ve_formul_yok():
    veri = excel.xlsx([["id", "sayi", "tarih", "metin", "bos"], ["A-1", "12", "1946-02-16", "=1+1", ""],
                       ["B-2", "007", "-3", "Güvenlik Konseyi", "x"]], "deneme")
    ws = load_workbook(io.BytesIO(veri)).active
    assert ws["B2"].value == 12 and ws["C2"].value == "1946-02-16" and ws["D2"].value == "=1+1" and ws["D2"].data_type == "s"
    assert ws["E2"].value is None and ws["B3"].value == "007" and ws["C3"].value == -3 and ws["D3"].value == "Güvenlik Konseyi"
